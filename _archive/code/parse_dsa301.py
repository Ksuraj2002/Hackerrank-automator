r"""Parse the DSA_301_questions tab of `Curriculum (1).xlsx` into a normalised
per-lecture question list.

Sheet schema: rows of (Lecture, Topic, Question, Difficulty), with a leading
"Classwork" banner row and, further down, a "Homework" banner. This script
walks both sections and tags each record with its kind.

Output: tools/out/dsa301_questions.json
"""
import argparse
import json
import os
import re

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "tools", "out")

DEFAULT_XLSX = r"C:\Users\sunst\Downloads\Curriculum (1).xlsx"
DEFAULT_SHEET = "DSA_301_questions"

LC_RE = re.compile(r"\bLC\s*[- ]?(\d{1,4})\b", re.IGNORECASE)
GFG_RE = re.compile(r"\(gfg\)", re.IGNORECASE)
LEC_RE = re.compile(r"^LEC[-\s]?(\d{1,3})$", re.IGNORECASE)


def slug(text):
    t = text.lower()
    t = re.sub(r"\(gfg\)|\(intro\)|\(review\)|\(lc\d+\)", " ", t)
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t[:80]


def canonical_key(text):
    m = LC_RE.search(text)
    if m:
        return "lc%d" % int(m.group(1))
    return slug(text)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--max-lec", type=int, default=None,
                    help="Cap at LEC-N inclusive (e.g. --max-lec 10)")
    ap.add_argument("--kind", choices=["Classwork", "Homework", "both"], default="both")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[args.sheet]
    print("Reading %s :: %s" % (os.path.basename(args.xlsx), ws.title))

    records = []
    kind = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        lec, topic, question, diff = (row + (None,) * 4)[:4]
        if isinstance(lec, str) and lec.strip() in ("Classwork", "Homework"):
            kind = lec.strip()
            continue
        if not lec or not question:
            continue
        if not isinstance(lec, str):
            continue
        m = LEC_RE.match(lec.strip())
        if not m:
            # header row "Lecture / Topic / Question / Difficulty"
            continue
        n = int(m.group(1))
        if args.max_lec and n > args.max_lec:
            continue
        if args.kind != "both" and kind != args.kind:
            continue
        text = str(question).strip()
        records.append({
            "kind": kind or "Classwork",
            "lecture": "LEC-%d" % n,
            "lec_num": n,
            "topic": (topic or "").strip(),
            "text": text,
            "difficulty": (diff or "").strip(),
            "lc": (lambda mm: int(mm.group(1)) if mm else None)(LC_RE.search(text)),
            "gfg": bool(GFG_RE.search(text)),
            "key": canonical_key(text),
        })

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "dsa301_questions.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)

    print("parsed %d items -> %s" % (len(records), out))
    by_lec = {}
    for r in records:
        by_lec.setdefault(r["lecture"], []).append(r)
    for lec in sorted(by_lec, key=lambda k: int(k.split("-")[1])):
        items = by_lec[lec]
        print("  %-7s %s (%d)" % (lec, items[0]["topic"], len(items)))


if __name__ == "__main__":
    main()
