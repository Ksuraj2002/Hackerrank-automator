"""Turn a curriculum sheet into two editable markdown challenge lists.

Reads one sheet of DSA_101_Curriculum.xlsx -- `DSA_101` or `DSA_301`, chosen with
--course -- and emits, for course 101:

    classwork.md   -- every Classwork problem, one row per problem
    homework.md    -- every Homework  problem, one row per problem

and for any other course the same two files under a course prefix
(`301-classwork.md`, `301-homework.md`). Which sheet, which filenames and which
lectures a course covers are defined in courses.py.

Each row carries a stable id (CW-D07-2 / HW-D07-2), the problem text exactly as
written in the sheet, a Type column, and a Slug column for you to fill in with
the HackerRank challenge slug. create_contests.py reads these files back.

Re-running this script regenerates the files and would discard slugs you have
typed in, so it refuses to overwrite unless you ask:

    python extract_curriculum.py                          # first run, DSA 101
    python extract_curriculum.py --course 301              # first run, DSA 301
    python extract_curriculum.py --merge                   # sheet changed, keep slugs
    python extract_curriculum.py --force                   # regenerate from scratch

A course may cover only part of its sheet: DSA 301 is limited to LEC-1..LEC-10,
so rows outside that range are read and then skipped. Widen the range in
courses.py and re-run with --merge to extend the files without losing slugs.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

import courses
import md_table

HERE = Path(__file__).resolve().parent
DEFAULT_XLSX = Path.home() / "Downloads" / "DSA_101_Curriculum.xlsx"

# Columns in the sheet (1-indexed): A Week, B Lecture, C Lecture Name,
# D Content, E Classwork, F Homework, G Status, H Lesson Plan, I Questions Prepared
COL_WEEK, COL_LECTURE, COL_NAME, COL_CLASSWORK, COL_HOMEWORK = 1, 2, 3, 5, 6

# Items that describe a paper exercise or a reading task rather than a program
# you can submit to a judge. These are marked Type=theory and skipped when the
# contests are built.
THEORY_PATTERNS = [
    r"\btc of\b",
    r"\bcalculate tc\b",
    r"\bdetermine tc\b",
    r"space complexity of\b",
    r"\bcalculate sc\b",
    r"\bread up on\b",
    r"\btrace\b",
    r"on paper\b",
    r"\bdry run\b",
    r"\bdifferentiate\b",
    r"compare stability\b",
    r"trade-off examples\b",
    r"truth tables\b",
    r"find o\(1\) space solutions\b",
    r"then verify\b",
    r"\bmock test\b",
    r"complete any pending\b",
    r"rest and prepare\b",
]
THEORY_RE = re.compile("|".join(THEORY_PATTERNS), re.IGNORECASE)

ITEM_RE = re.compile(r"^\s*(\d+)[.)]\s*(.+?)\s*$")


def split_items(cell):
    """Turn a '1. foo\\n2. bar' cell into ['foo', 'bar'].

    A line with no leading number is treated as a continuation of the previous
    item, so a wrapped problem statement stays in one piece.
    """
    if not cell:
        return []
    items = []
    for raw in str(cell).replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        m = ITEM_RE.match(line)
        if m:
            items.append(m.group(2))
        elif items:
            items[-1] += " " + line
        else:
            items.append(line)
    return [i.strip() for i in items if i.strip()]


def day_number(lecture):
    m = re.search(r"(\d+)", str(lecture))
    return int(m.group(1)) if m else 0


def read_sheet(xlsx, sheet, keep=None):
    """Read one sheet. `keep` is a predicate on the lecture number, or None."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit("sheet %r not found in %s (have: %s)" % (sheet, xlsx, wb.sheetnames))
    ws = wb[sheet]

    days, skipped = [], 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        lecture = row[COL_LECTURE - 1]
        if not lecture:
            continue
        if keep and not keep(day_number(lecture)):
            skipped += 1
            continue
        week = row[COL_WEEK - 1]
        days.append(
            {
                "day": day_number(lecture),
                "lecture": str(lecture).strip(),
                "week": int(week) if isinstance(week, (int, float)) else 0,
                "topic": str(row[COL_NAME - 1] or "").strip(),
                "classwork": split_items(row[COL_CLASSWORK - 1]),
                "homework": split_items(row[COL_HOMEWORK - 1]),
            }
        )
    return days, skipped


# --- writing ---------------------------------------------------------------

HEADER = """# {course_title} - {kind_title}

Generated from `DSA_101_Curriculum.xlsx`, sheet `{sheet}`, column `{kind_title}`.
This file is the source of truth for `create_contests.py` - edit it here, not in
the spreadsheet, once you start filling in slugs.

- **Contest name:** {contest_name}
- **Contest slug:** `{contest_slug}`
- **Coverage:** {scope}.
- **Counts:** {n_total} problems total - {n_code} codeable, {n_theory} theory.

## How to fill this in

Every problem needs a **Slug** before it can be added to the contest. The slug is
the last path segment of the challenge URL on HackerRank:

```
https://www.hackerrank.com/challenges/simple-array-sum/problem
                                      ^^^^^^^^^^^^^^^^ this part
```

- Put the bare slug in the Slug column. No full URL, no backticks needed.
- Leave Slug empty for anything not mapped yet. Those are reported as unmapped
  and skipped; everything else still gets added.
- Rows with Type `theory` are paper or reading exercises with no judge problem.
  They are always skipped, whatever the Slug column says.
- `python resolve_slugs.py --course {course} --{kind}` uses this account's own
  challenge bank plus HackerRank's search to propose candidate slugs and writes
  them into this file. Review them before trusting them - a keyword search will
  happily return the wrong problem.
- Order matters: challenges are added to the contest in the order they appear
  below, so the first lecture lands first.

---
"""


def write_markdown(path, kind, days, carried, course_id, spec):
    course = courses.get(course_id)
    kind_title = "Classwork" if kind == "classwork" else "Homework"
    prefix = "CW" if kind == "classwork" else "HW"

    rows_by_week = {}
    n_total = n_code = n_theory = 0

    for d in days:
        entries = []
        for idx, problem in enumerate(d[kind], start=1):
            pid = "%s-D%02d-%d" % (prefix, d["day"], idx)
            ptype = "theory" if THEORY_RE.search(problem) else "code"
            entries.append((pid, problem, ptype, carried.get(pid, "")))
            n_total += 1
            if ptype == "theory":
                n_theory += 1
            else:
                n_code += 1
        if entries:
            rows_by_week.setdefault(d["week"], []).append((d, entries))

    out = [
        HEADER.format(
            course=course_id,
            course_title=course["title"],
            sheet=course["sheet"],
            scope=course["scope"],
            kind=kind,
            kind_title=kind_title,
            contest_name=spec["name"],
            contest_slug=spec["slug"],
            n_total=n_total,
            n_code=n_code,
            n_theory=n_theory,
        )
    ]

    for week in sorted(rows_by_week):
        out.append("\n## Week %d\n" % week)
        for d, entries in rows_by_week[week]:
            out.append("### %s - %s\n" % (d["lecture"], d["topic"]))
            out.append("| ID | Problem | Slug | Type |")
            out.append("|----|---------|------|------|")
            for pid, problem, ptype, slug in entries:
                out.append(md_table.format_row(pid, problem, slug, ptype))
            out.append("")

    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    print("wrote %s: %d problems (%d code, %d theory)" % (path.name, n_total, n_code, n_theory))


def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    courses.add_argument(ap)
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help="path to the workbook")
    ap.add_argument("--out-dir", type=Path, default=HERE)
    ap.add_argument("--force", action="store_true", help="overwrite markdown, discarding slugs")
    ap.add_argument("--merge", action="store_true", help="overwrite markdown but keep slugs")
    args = ap.parse_args()

    if not args.xlsx.exists():
        sys.exit("workbook not found: %s" % args.xlsx)

    course = courses.get(args.course)
    keep = lambda n: courses.in_range(args.course, n)
    days, skipped = read_sheet(args.xlsx, course["sheet"], keep)
    print("read %d lecture rows from %s, sheet %s"
          % (len(days), args.xlsx.name, course["sheet"]))
    if skipped:
        print("  skipped %d row(s) outside %s" % (skipped, course["scope"]))

    for spec in courses.contests(args.course):
        kind = spec["kind"]
        path = args.out_dir / spec["source"]
        if path.exists() and not (args.force or args.merge):
            print("skip %s: already exists (--merge keeps slugs, --force wipes them)" % path.name)
            continue
        carried = md_table.existing_slugs(path) if args.merge else {}
        if carried:
            print("  carrying %d slug(s) forward into %s" % (len(carried), path.name))
        write_markdown(path, kind, days, carried, args.course, spec)


if __name__ == "__main__":
    main()
