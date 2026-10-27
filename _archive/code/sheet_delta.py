r"""What is new in the standardised sheet compared with the previous curriculum,
and does each of those questions already exist on HackerRank?

  venv\Scripts\python.exe tools\sheet_delta.py

Writes tools/out/new_questions.json  -> the additional questions, each tagged
                                        exists / review / absent
       tools/out/new_questions.txt   -> the same, readable

Feed the absent ones straight into the builder:
  venv\Scripts\python.exe tools\build_md.py --from tools\out\new_questions.json
"""
import argparse
import importlib.util
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOOLS = os.path.join(ROOT, "tools")
OUT_DIR = os.path.join(TOOLS, "out")
sys.path.insert(0, TOOLS)

import diff_challenges as dc  # noqa: E402
import parse_curriculum as pc  # noqa: E402

DEFAULT_OLD = r"C:\Users\sunst\Downloads\DSA_101_Curriculum.xlsx"


def load_sheet(path, sheet=None):
    """Parse one curriculum workbook into codeable question records."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    out, week = [], None
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=7):
        wk, day, name, _content, classwork, homework, _status = [c.value for c in row]
        if wk is not None:
            week = int(wk)
        if not day:
            continue
        for kind, cell in (("Classwork", classwork), ("Homework", homework)):
            for i, text in enumerate(pc.split_items(cell), start=1):
                if len(text) < 3 or not dc.is_challenge(text):
                    continue
                out.append({"week": week, "day": str(day).strip(), "kind": kind, "index": i,
                            "lecture_name": (name or "").strip().replace("\n", " "),
                            "text": text, "lc": None, "key": pc.canonical_key(text)})
    return out


def hr_index():
    path = os.path.join(OUT_DIR, "hackerrank_challenges.json")
    if not os.path.exists(path):
        raise SystemExit("Missing %s - run tools/hr_list_challenges.py first." % path)
    with open(path, encoding="utf-8") as f:
        dump = json.load(f)
    index, seen = [], set()
    for rec in dump["records"]:
        name = (rec.get("name") or "").strip()
        if not name or name in seen:
            continue
        seen.add(name)
        index.append((dc.norm(name), {"name": name, "slug": rec.get("slug", ""),
                                      "id": rec.get("id")}))
    return index


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--old", default=DEFAULT_OLD, help="previous curriculum workbook")
    ap.add_argument("--old-sheet", default="DSA_101")
    ap.add_argument("--new", default=pc.DEFAULT_XLSX, help="standardised workbook")
    ap.add_argument("--new-sheet", default=None)
    args = ap.parse_args()

    old = load_sheet(args.old, args.old_sheet)
    new = load_sheet(args.new, args.new_sheet)
    old_keys = {r["key"] for r in old}

    # One record per question that the standardised sheet added.
    added, seen = [], set()
    for rec in new:
        if rec["key"] in old_keys or rec["key"] in seen:
            continue
        seen.add(rec["key"])
        added.append(rec)

    verdicts = {}
    verdict_path = os.path.join(TOOLS, "verdicts.json")
    if os.path.exists(verdict_path):
        with open(verdict_path, encoding="utf-8") as f:
            verdicts = {k: v for k, v in json.load(f).items() if not k.startswith("_")}

    index = hr_index()
    for rec in added:
        cand, sc, drift = dc.best_candidate(dc.norm(rec["text"]), index)
        rec["hr_candidate"] = cand["name"] if cand else ""
        rec["hr_slug"] = cand["slug"] if cand else ""
        rec["hr_score"] = round(sc, 3)
        rec["status"] = ("exists" if sc >= dc.AUTO_MATCH
                         else "review" if sc >= dc.REVIEW_MIN else "absent")
        rec["modifier_drift"] = drift

        # A recorded human call always wins over the score.
        verdict = verdicts.get(rec["key"])
        if verdict:
            rec["status"] = verdict["status"]
            rec["verdict_note"] = verdict.get("note", "")
            if verdict.get("hr"):
                rec["hr_candidate"] = verdict["hr"]

    with open(os.path.join(OUT_DIR, "new_questions.json"), "w", encoding="utf-8") as f:
        json.dump(added, f, indent=2, ensure_ascii=False)

    order = {"create": 0, "absent": 1, "review": 2, "exists": 3, "drill": 4}
    with open(os.path.join(OUT_DIR, "new_questions.txt"), "w", encoding="utf-8") as f:
        f.write("# Questions the standardised sheet adds, vs %s [%s]\n\n"
                % (os.path.basename(args.old), args.old_sheet))
        for rec in sorted(added, key=lambda r: (order[r["status"]], r["day"])):
            f.write("%-7s %-7s %-9s %s\n"
                    % (rec["status"].upper(), rec["day"], rec["kind"], rec["text"][:96]))
            if rec["hr_candidate"] and rec["status"] in ("exists", "review"):
                f.write("        existing: %s\n" % rec["hr_candidate"])
            if rec.get("verdict_note"):
                f.write("        note: %s\n" % rec["verdict_note"])
        f.write("\n")

    counts = {}
    for rec in added:
        counts[rec["status"]] = counts.get(rec["status"], 0) + 1
    print("old sheet: %d codeable items (%d unique)" % (len(old), len(old_keys)))
    print("new sheet: %d codeable items" % len(new))
    print("questions the standardised sheet ADDS: %d unique" % len(added))
    for status in ("create", "absent", "review", "exists", "drill"):
        if counts.get(status):
            print("  %-8s %d" % (status, counts[status]))
    print("\n-> tools/out/new_questions.txt")


if __name__ == "__main__":
    main()
